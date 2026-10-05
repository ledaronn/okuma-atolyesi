"""Belge editörü penceresi (ekran dışı Qt): yaz, biçimle, kaydet, PDF, kayıp koruması."""
import os
import sys
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import fitz
import pytest
from PySide6.QtGui import QFont, QTextCursor, QTextFormat, QTextListFormat
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

from belge import editor as ed


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def pencere(app, monkeypatch):
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Discard)
    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: None)
    w = ed.BelgeEditoru()
    w.resize(1200, 860); w.show(); app.processEvents()
    yield w
    w.gorunum.document().setModified(False)
    w.close()


def yaz(w, metin):
    w.gorunum.textCursor().insertText(metin)


def test_yaz_bicimle_kaydet_ac(pencere, tmp_path, monkeypatch):
    w = pencere
    w.stil.setCurrentIndex(2); w.stil_uygula(2)                 # Baslik 1
    yaz(w, "Rapor Başlığı")
    w.gorunum.textCursor().insertBlock()
    w.stil_uygula(0)
    w.kalin.setChecked(True); w.kalin_yap(); yaz(w, "kalın metin"); w.kalin.setChecked(False); w.kalin_yap()
    w.gorunum.textCursor().insertBlock()
    w.liste(QTextListFormat.ListDecimal); yaz(w, "birinci")
    w.gorunum.textCursor().insertBlock(); yaz(w, "ikinci")
    w.liste(QTextListFormat.ListDecimal)                        # ayni dugme listeden cikarir
    yaz(w, " (liste dışı)")
    yol = tmp_path / "rapor.docx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(yol), ""))
    assert w.kaydet() and yol.exists() and not w.gorunum.document().isModified()
    assert w.windowTitle().startswith("rapor.docx")
    w2 = ed.BelgeEditoru(str(yol))
    b = w2.gorunum.document()
    bloklar = []
    blk = b.begin()
    while blk.isValid():
        bloklar.append(blk); blk = blk.next()
    assert bloklar[0].text() == "Rapor Başlığı" and bloklar[0].blockFormat().headingLevel() == 1
    assert bloklar[2].textList() is not None and bloklar[2].text() == "birinci"
    assert bloklar[3].textList() is None and bloklar[3].text() == "ikinci (liste dışı)"
    w2.gorunum.document().setModified(False); w2.close()


PDF_BETIK = r"""
import json, sys
sys.path.insert(0, sys.argv[1])
from PySide6.QtWidgets import QApplication
app = QApplication([])
from belge import editor as ed
w = ed.BelgeEditoru(); w.resize(1200, 860)
w.ayar.ust_bilgi = "Gizli Rapor"; w.ayar.sayfa_no = True; w._duzen_degisti()
for i in range(60):
    w.gorunum.textCursor().insertText(f"Paragraf {i}: " + "uzun bir cümle " * 8); w.gorunum.textCursor().insertBlock()
w.sayfa_sonu(); w.gorunum.textCursor().insertText("Son sayfa")
w.pdf_yaz(sys.argv[2])
print(json.dumps({"sayfa": w.gorunum.sayfa_sayisi()}))
w.gorunum.document().setModified(False)
"""


@pytest.mark.skipif(sys.platform != "win32", reason="PDF metni gercek platformda sinanir (offscreen yazi tiplerini sekil olarak gomer)")
def test_pdf_sayfa_ust_alt_bilgi(tmp_path):
    """Gercek Windows platformunda (offscreen yazi tiplerini sekil olarak gomer,
    metin okunamaz): PDF sayfa sayisi ekrandakiyle ayni, her sayfada ust bilgi
    ve sayfa numarasi SECILEBILIR metin, sayfa sonu dogru yerde."""
    import json
    import subprocess
    pdf = tmp_path / "cikti.pdf"
    ortam = dict(os.environ); ortam.pop("QT_QPA_PLATFORM", None)
    r = subprocess.run([sys.executable, "-c", PDF_BETIK, str(Path(__file__).resolve().parent.parent), str(pdf)],
                       capture_output=True, text=True, env=ortam, timeout=120)
    assert r.returncode == 0, r.stderr[-800:]
    beklenen = json.loads(r.stdout.strip().splitlines()[-1])["sayfa"]
    assert beklenen >= 3
    d = fitz.open(pdf)
    assert d.page_count == beklenen, "PDF ekrandaki sayfa sayisiyla ayni"
    assert round(d[0].rect.width / 72 * 25.4) == 210 and round(d[0].rect.height / 72 * 25.4) == 297
    for i, s in enumerate(d):
        metin = s.get_text()
        assert "Gizli Rapor" in metin and str(i + 1) in metin.split()
    assert "Son sayfa" in d[-1].get_text() and "Son sayfa" not in d[-2].get_text()


def test_stil_bos_blokta_yazilacak_metne_uygulanir(pencere):
    w = pencere
    w.stil_uygula(1)                                            # Belge basligi, bos ilk blok
    yaz(w, "Kompost Rehberi")
    f = w.gorunum.document().begin().begin().fragment().charFormat()
    assert f.fontPointSize() == 26 and f.fontWeight() >= QFont.Bold
    w.gorunum.textCursor().insertBlock(); w.stil_uygula(0); yaz(w, "normal")
    assert w.gorunum.document().lastBlock().begin().fragment().charFormat().fontWeight() < QFont.Bold


def test_sayfa_duzeni(pencere):
    w = pencere
    w.kenar.setCurrentIndex(1); w.kenar_sec(1)                  # Dar
    assert w.ayar.sol == 12.7
    w.yon_degistir()
    assert w.ayar.yatay and round(w.gorunum.document().pageSize().width()) == round(297 * ed.MM_PX)


def test_tablo_ve_bul_degistir(pencere, monkeypatch):
    w = pencere
    yaz(w, "elma armut elma")
    w.gorunum.textCursor().insertBlock()
    from PySide6.QtGui import QTextLength, QTextTableFormat
    t = w.gorunum.textCursor().insertTable(2, 2, QTextTableFormat())
    w.gorunum.setTextCursor(t.cellAt(0, 0).firstCursorPosition())
    w._tablo_islem("satir"); w._tablo_islem("sutun")
    assert (t.rows(), t.columns()) == (3, 3)
    w.bul_ac(); w._bul.bul.setText("elma"); w._bul.yeni.setText("ayva"); w._bul.tumu()
    assert "ayva armut ayva" in w.gorunum.toPlainText() and "2 yer" in w._bul.durum.text()
    w._bul.close()


def test_word_belgesi_ustune_kaydedilmez(pencere, tmp_path, monkeypatch):
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = docx.Document(); p = d.add_paragraph("Word belgesi")
    ref = OxmlElement("w:commentReference"); ref.set(qn("w:id"), "0")
    r = OxmlElement("w:r"); r.append(ref); p._p.append(r)
    ozgun = tmp_path / "word.docx"; d.save(str(ozgun))
    once = ozgun.read_bytes()
    w = pencere
    assert w.ac(str(ozgun))
    assert w.uyari.isVisible() and "yorumlar" in w.uyari_metin.text()
    yaz(w, " değişti")
    # kaydet -> farkli kaydet; kullanici yine ozgun dosyayi secerse reddedilir
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(ozgun), ""))
    assert w.kaydet() is False
    assert ozgun.read_bytes() == once, "ozgun Word dosyasi degismedi"
    yeni = tmp_path / "word (düzenlendi).docx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(yeni), ""))
    assert w.kaydet() and yeni.exists() and ozgun.read_bytes() == once


def test_libreoffice_ozgun_dosyayi_acar(pencere, tmp_path, monkeypatch):
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = docx.Document(); p = d.add_paragraph("izli")
    ins = OxmlElement("w:ins"); ins.set(qn("w:id"), "1"); ins.set(qn("w:author"), "A"); p._p.append(ins)
    ozgun = tmp_path / "izli.docx"; d.save(str(ozgun))
    cagri = []
    monkeypatch.setattr(ed, "soffice_yolu", lambda: "soffice.exe")
    monkeypatch.setattr(ed.subprocess, "Popen", lambda args, **k: cagri.append(args))
    w = pencere
    w.ac(str(ozgun)); yaz(w, "x")
    w.libreoffice_ac()
    assert cagri == [["soffice.exe", "--writer", str(ozgun)]], "kayipli belgede OZGUN dosya acilir, kaydetmeden"
