"""Belge editörünün .docx çekirdeği: sıfırdan yaz -> kaydet -> geri oku, kayıp koruması."""
import os
import sys
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest
from PySide6.QtCore import Qt, QUrl
from PySide6.QtGui import (QBrush, QColor, QFont, QGuiApplication, QImage, QTextBlockFormat, QTextCharFormat,
                           QTextCursor, QTextDocument, QTextFormat, QTextImageFormat, QTextListFormat)

from belge import docx_io
from belge.docx_io import SayfaAyari, docx_oku, docx_yaz


@pytest.fixture(scope="module", autouse=True)
def uygulama():
    # QApplication (QGuiApplication degil): ayni pytest surecinde sonra pencere
    # acan testler QGuiApplication ile cokuyordu.
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def ornek_belge():
    """Faz 1'in her özelliğinden bir parça içeren belge."""
    b = QTextDocument()
    c = QTextCursor(b)
    bf = QTextBlockFormat(); bf.setHeadingLevel(1)
    c.setBlockFormat(bf); c.insertText("Ana Başlık", docx_io.baslik_bicimi(1))
    # bicimli paragraf
    c.insertBlock(QTextBlockFormat(), QTextCharFormat())
    k = QTextCharFormat(); k.setFontWeight(QFont.Bold); c.insertText("kalın ", k)
    i = QTextCharFormat(); i.setFontItalic(True); i.setFontUnderline(True); c.insertText("italik-altçizgi ", i)
    r = QTextCharFormat(); r.setForeground(QBrush(QColor("#c00000"))); r.setFontPointSize(14)
    r.setFontFamilies(["Georgia"]); c.insertText("kırmızı ", r)
    v = QTextCharFormat(); v.setBackground(QBrush(QColor("#ffff00"))); c.insertText("vurgulu ", v)
    u = QTextCharFormat(); u.setVerticalAlignment(QTextCharFormat.AlignSuperScript); c.insertText("2", u)
    s = QTextCharFormat(); s.setFontStrikeOut(True); c.insertText(" üstü çizili", s)
    # ortali, 1.5 satir, sayfa sonu oncesi
    of = QTextBlockFormat(); of.setAlignment(Qt.AlignHCenter)
    of.setLineHeight(150, QTextBlockFormat.ProportionalHeight.value)
    c.insertBlock(of, QTextCharFormat()); c.insertText("Ortalı satır")
    sf = QTextBlockFormat(); sf.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysBefore)
    c.insertBlock(sf, QTextCharFormat()); c.insertText("Yeni sayfa")
    # madde listesi + numarali liste
    c.insertBlock(QTextBlockFormat(), QTextCharFormat())
    lf = QTextListFormat(); lf.setStyle(QTextListFormat.ListDisc)
    c.createList(lf); c.insertText("madde bir")
    c.insertBlock(); c.insertText("madde iki")
    c.insertBlock(QTextBlockFormat(), QTextCharFormat())
    nf = QTextListFormat(); nf.setStyle(QTextListFormat.ListDecimal)
    c.createList(nf); c.insertText("numara bir")
    # tablo
    c.insertBlock(QTextBlockFormat(), QTextCharFormat())
    t = c.insertTable(2, 3)
    for ri in range(2):
        for ci in range(3):
            t.cellAt(ri, ci).firstCursorPosition().insertText(f"h{ri}{ci}")
    t.mergeCells(1, 0, 1, 2)
    c.movePosition(QTextCursor.End)
    # resim
    img = QImage(40, 20, QImage.Format_RGB32); img.fill(QColor("#336699"))
    b.addResource(QTextDocument.ImageResource, QUrl("belge://test"), img)
    imf = QTextImageFormat(); imf.setName("belge://test"); imf.setWidth(40); imf.setHeight(20)
    c.insertBlock(QTextBlockFormat(), QTextCharFormat()); c.insertImage(imf)
    return b


def bloklar(b):
    blk = b.begin()
    while blk.isValid():
        yield blk
        blk = blk.next()


def parca_bicimi(b, metin):
    for blk in bloklar(b):
        it = blk.begin()
        while not it.atEnd():
            f = it.fragment()
            if f.isValid() and metin in f.text():
                return f.charFormat()
            it += 1
    raise AssertionError(f"parca yok: {metin}")


def test_gidis_donus(tmp_path):
    yol = tmp_path / "deneme.docx"
    sayfa = SayfaAyari(ust=20, alt=22, sol=30, sag=15, ust_bilgi="Rapor", alt_bilgi="Okuma Atölyesi", sayfa_no=True)
    docx_yaz(ornek_belge(), sayfa, yol)
    sonuc = docx_oku(yol)
    b = sonuc.belge
    assert sonuc.uyarilar == []
    metinler = [blk.text() for blk in bloklar(b)]
    assert "Ana Başlık" in metinler
    baslik = next(blk for blk in bloklar(b) if blk.text() == "Ana Başlık")
    assert baslik.blockFormat().headingLevel() == 1
    assert parca_bicimi(b, "kalın").fontWeight() >= QFont.Bold
    ital = parca_bicimi(b, "italik-altçizgi")
    assert ital.fontItalic() and ital.fontUnderline()
    kir = parca_bicimi(b, "kırmızı")
    assert kir.foreground().color().name() == "#c00000" and kir.fontPointSize() == 14 and kir.fontFamilies()[0] == "Georgia"
    assert parca_bicimi(b, "vurgulu").background().color().name() == "#ffff00"
    assert parca_bicimi(b, "2").verticalAlignment() == QTextCharFormat.AlignSuperScript
    assert parca_bicimi(b, "üstü çizili").fontStrikeOut()
    orta = next(blk for blk in bloklar(b) if blk.text() == "Ortalı satır")
    assert orta.blockFormat().alignment() & Qt.AlignHCenter
    assert abs(orta.blockFormat().lineHeight() - 150) < 1
    yeni = next(blk for blk in bloklar(b) if blk.text() == "Yeni sayfa")
    assert yeni.blockFormat().pageBreakPolicy() & QTextFormat.PageBreak_AlwaysBefore
    madde = [blk for blk in bloklar(b) if blk.text().startswith("madde")]
    assert len(madde) == 2 and madde[0].textList() is not None and madde[0].textList() == madde[1].textList()
    assert madde[0].textList().format().style() == QTextListFormat.ListDisc
    numara = next(blk for blk in bloklar(b) if blk.text() == "numara bir")
    assert numara.textList().format().style() == QTextListFormat.ListDecimal
    # tablo: 2x3, alt satirda yatay birlestirme
    from PySide6.QtGui import QTextTable
    tablolar = [f for f in b.rootFrame().childFrames() if isinstance(f, QTextTable)]
    assert len(tablolar) == 1 and tablolar[0].rows() == 2 and tablolar[0].columns() == 3
    assert tablolar[0].cellAt(0, 2).firstCursorPosition().block().text() == "h02"
    assert tablolar[0].cellAt(1, 0).columnSpan() == 2
    # resim
    resimler = [f for blk in bloklar(b) for f in _parcalar(blk) if f.charFormat().isImageFormat()]
    assert len(resimler) == 1 and round(resimler[0].charFormat().toImageFormat().width()) == 40
    # sayfa duzeni
    s = sonuc.sayfa
    assert (round(s.ust), round(s.alt), round(s.sol), round(s.sag)) == (20, 22, 30, 15)
    assert round(s.genislik) == 210 and round(s.yukseklik) == 297
    assert s.ust_bilgi == "Rapor" and "Okuma Atölyesi" in s.alt_bilgi and s.sayfa_no


def _parcalar(blk):
    it = blk.begin()
    while not it.atEnd():
        f = it.fragment()
        if f.isValid():
            yield f
        it += 1


def test_yatay_sayfa(tmp_path):
    yol = tmp_path / "yatay.docx"
    s = SayfaAyari(); s.dondur()
    b = QTextDocument(); QTextCursor(b).insertText("yatay")
    docx_yaz(b, s, yol)
    o = docx_oku(yol).sayfa
    assert o.yatay and round(o.genislik) == 297


def test_word_belgesinde_desteksiz_icerik_uyarilir(tmp_path):
    """Word'de yazılmış, izlenen değişiklik ve yorum içeren belge: uyarı listesi dolu."""
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = docx.Document()
    p = d.add_paragraph("Önce")
    ins = OxmlElement("w:ins"); ins.set(qn("w:id"), "1"); ins.set(qn("w:author"), "A")
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "eklenen"; r.append(t); ins.append(r)
    p._p.append(ins)
    ref = OxmlElement("w:commentReference"); ref.set(qn("w:id"), "0")
    rr = OxmlElement("w:r"); rr.append(ref); p._p.append(rr)
    yol = tmp_path / "word.docx"; d.save(str(yol))
    uyarilar = docx_oku(yol).uyarilar
    assert "izlenen değişiklikler" in uyarilar and "yorumlar" in uyarilar


def test_bos_belge(tmp_path):
    yol = tmp_path / "bos.docx"
    docx_yaz(QTextDocument(), SayfaAyari(), yol)
    assert docx_oku(yol).belge.toPlainText() == ""
