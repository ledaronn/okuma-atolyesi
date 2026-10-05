"""belge/docx_io.py — .docx <-> QTextDocument donusumu.

Editor belgeyi bellekte QTextDocument olarak tutar (Qt'nin metin motoru:
bicim, liste, tablo, resim, sayfalama). Diske .docx (python-docx) yazilir.

Desteklenen (Faz 1): paragraflar ve run bicimi (yazi tipi, boyut, kalin,
italik, alti/ustu cizili, renk, vurgu, ust/alt simge), baslik stilleri,
hizalama, satir araligi, paragraf oncesi/sonrasi bosluk, girinti, madde ve
numara listeleri (iki duzey), sayfa sonu, tablolar (yatay birlestirme), satir
ici resimler, sayfa boyutu/yonu/kenar bosluklari, ust/alt bilgi metni ve
sayfa numarasi.

KAYIP KORUMASI: Word'de yazilmis bir dosyada editorun TASIYAMADIGI bir sey
varsa (izlenen degisiklik, yorum, dipnot, metin kutusu, kayan resim, icerik
denetimi, denklem, grafik, birden cok bolum...) `docx_oku` bunu `uyarilar`
listesinde soyler. Editor bu durumda ozgun dosyanin USTUNE kaydetmez: o
icerik sessizce silinirdi.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from PySide6.QtCore import QBuffer, QByteArray, QIODevice, QUrl
from PySide6.QtGui import (QBrush, QColor, QFont, QImage, QTextBlockFormat, QTextCharFormat, QTextCursor,
                           QTextDocument, QTextFormat, QTextFrame, QTextImageFormat, QTextLength, QTextListFormat,
                           QTextTableFormat)

EKRAN_DPI = 96.0
MM_PX = EKRAN_DPI / 25.4          # 1 mm = 3.78 px (96 dpi)
PT_PX = EKRAN_DPI / 72.0
EMU_INC = 914400


@dataclass
class SayfaAyari:
    """Belgenin sayfa duzeni (mm). Tek bolum: Faz 1."""
    genislik: float = 210.0
    yukseklik: float = 297.0
    ust: float = 25.0
    alt: float = 25.0
    sol: float = 25.0
    sag: float = 25.0
    ust_bilgi: str = ""
    alt_bilgi: str = ""
    sayfa_no: bool = False          # alt bilgide "sayfa n" alani

    @property
    def yatay(self) -> bool:
        return self.genislik > self.yukseklik

    def dondur(self) -> None:
        self.genislik, self.yukseklik = self.yukseklik, self.genislik


@dataclass
class OkumaSonucu:
    belge: QTextDocument
    sayfa: SayfaAyari
    uyarilar: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------
# Kayip korumasi: editorun tasiyamadigi Word icerigi
# --------------------------------------------------------------------------
DESTEKSIZ = [
    (r"<w:ins\b|<w:del\b", "izlenen değişiklikler"),
    (r"<w:commentReference\b", "yorumlar"),
    (r"<w:footnoteReference\b|<w:endnoteReference\b", "dipnot/son notlar"),
    (r"<w:txbxContent\b", "metin kutuları"),
    (r"<wp:anchor\b", "kayan (metin dışı) resim/şekil"),
    (r"<w:sdt\b", "içerik denetimi / içindekiler alanı"),
    (r"<m:oMath\b", "denklemler"),
    (r"<c:chart\b", "grafikler"),
    (r'<w:vMerge\b', "dikey birleştirilmiş tablo hücreleri"),
]


def desteksiz_icerik(belge_xml: str, bolum_sayisi: int) -> list[str]:
    bulunan = [ad for kalip, ad in DESTEKSIZ if re.search(kalip, belge_xml)]
    if bolum_sayisi > 1:
        bulunan.append(f"birden çok bölüm ({bolum_sayisi}); yalnızca ilkinin sayfa düzeni korunur")
    return bulunan


# --------------------------------------------------------------------------
# Ortak eslemeler
# --------------------------------------------------------------------------
BASLIK_STILI = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}
BASLIK_BOYUT = {0: 26, 1: 20, 2: 16, 3: 13}      # 0 = belge basligi (Title)
HIZA_QT = {"LEFT": 0x1, "RIGHT": 0x2, "CENTER": 0x4, "JUSTIFY": 0x8}

# Word'un vurgu renkleri (WD_COLOR_INDEX) — vurgu bu 15 renkten biridir
VURGU = {"YELLOW": "#ffff00", "BRIGHT_GREEN": "#00ff00", "TURQUOISE": "#00ffff", "PINK": "#ff00ff",
         "BLUE": "#0000ff", "RED": "#ff0000", "DARK_BLUE": "#000080", "TEAL": "#008080", "GREEN": "#008000",
         "VIOLET": "#800080", "DARK_RED": "#800000", "DARK_YELLOW": "#808000", "GRAY_50": "#808080",
         "GRAY_25": "#c0c0c0", "BLACK": "#000000"}


def en_yakin_vurgu(renk: QColor) -> str:
    def uz(h):
        c = QColor(h)
        return (c.red() - renk.red()) ** 2 + (c.green() - renk.green()) ** 2 + (c.blue() - renk.blue()) ** 2
    return min(VURGU, key=lambda k: uz(VURGU[k]))


def baslik_bicimi(seviye: int) -> QTextCharFormat:
    f = QTextCharFormat()
    f.setFontWeight(QFont.Bold)
    f.setFontPointSize(BASLIK_BOYUT.get(seviye, 13))
    return f


# --------------------------------------------------------------------------
# OKUMA: .docx -> QTextDocument
# --------------------------------------------------------------------------
def docx_oku(yol) -> OkumaSonucu:
    import docx
    from docx.oxml.ns import qn
    d = docx.Document(str(yol))
    belge = QTextDocument()
    belge.setDefaultFont(_varsayilan_yazi(d))
    imlec = QTextCursor(belge)
    sayfa = _sayfa_oku(d)
    ilk = [True]
    resim_no = [0]
    numara = _numaralandirma(d)

    def yeni_blok(bfmt: QTextBlockFormat, cfmt: QTextCharFormat | None = None):
        if ilk[0]:
            imlec.setBlockFormat(bfmt)
            if cfmt is not None:
                imlec.setBlockCharFormat(cfmt)
            ilk[0] = False
        else:
            imlec.insertBlock(bfmt, cfmt or QTextCharFormat())

    sonraki_sayfa_sonu = [False]
    for oge in d.element.body.iterchildren():
        if oge.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            p = Paragraph(oge, d)
            bfmt, liste, cfmt0 = _paragraf_bicimi(p, numara)
            if sonraki_sayfa_sonu[0]:
                bfmt.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysBefore)
                sonraki_sayfa_sonu[0] = False
            yeni_blok(bfmt, cfmt0)
            if liste is not None:
                _listeye_kat(imlec, liste)
            sonraki_sayfa_sonu[0] = _runlari_yaz(imlec, p, d, belge, resim_no, cfmt0)
        elif oge.tag == qn("w:tbl"):
            from docx.table import Table
            t = Table(oge, d)
            if ilk[0]:
                ilk[0] = False
            _tablo_oku(imlec, t, d, belge, resim_no, numara)
            imlec.movePosition(QTextCursor.End)
    xml = d.element.xml
    uyarilar = desteksiz_icerik(xml, len(d.sections))
    belge.setModified(False)
    return OkumaSonucu(belge, sayfa, uyarilar)


def _varsayilan_yazi(d) -> QFont:
    stil = d.styles["Normal"].font
    f = QFont(stil.name or "Calibri")
    f.setPointSizeF(stil.size.pt if stil.size else 11.0)
    return f


def _sayfa_oku(d) -> SayfaAyari:
    s = d.sections[0]
    mm = lambda v, var: (v.mm if v is not None else var)   # noqa: E731
    ayar = SayfaAyari(genislik=mm(s.page_width, 210.0), yukseklik=mm(s.page_height, 297.0),
                      ust=mm(s.top_margin, 25.0), alt=mm(s.bottom_margin, 25.0),
                      sol=mm(s.left_margin, 25.0), sag=mm(s.right_margin, 25.0))
    try:
        ayar.ust_bilgi = "\n".join(p.text for p in s.header.paragraphs if p.text).strip()
        alt = s.footer
        ayar.alt_bilgi = "\n".join(p.text for p in alt.paragraphs if p.text).strip()
        ayar.sayfa_no = bool(re.search(r"\bPAGE\b", alt._element.xml))
    except Exception:
        pass
    return ayar


def _numaralandirma(d) -> dict:
    """numId -> {ilvl: 'bullet'|'decimal'...}. Liste turunu (madde/numara) belirlemek icin."""
    sonuc: dict = {}
    try:
        from docx.oxml.ns import qn
        kok = d.part.numbering_part.element
    except Exception:
        return sonuc
    soyut = {}
    for a in kok.findall(qn("w:abstractNum")):
        seviyeler = {}
        for lvl in a.findall(qn("w:lvl")):
            fmt = lvl.find(qn("w:numFmt"))
            seviyeler[int(lvl.get(qn("w:ilvl")))] = fmt.get(qn("w:val")) if fmt is not None else "decimal"
        soyut[a.get(qn("w:abstractNumId"))] = seviyeler
    for n in kok.findall(qn("w:num")):
        ref = n.find(qn("w:abstractNumId"))
        if ref is not None:
            sonuc[n.get(qn("w:numId"))] = soyut.get(ref.get(qn("w:val")), {})
    return sonuc


def _paragraf_bicimi(p, numara):
    """(blok bicimi, liste bilgisi ya da None, baslik karakter bicimi ya da None)."""
    from docx.oxml.ns import qn
    bfmt = QTextBlockFormat()
    stil = (p.style.name if p.style is not None else "") or ""
    cfmt = None
    m = re.match(r"^(Heading|Başlık)\s*(\d)$", stil)
    if m:
        seviye = min(int(m.group(2)), 6)
        bfmt.setHeadingLevel(seviye)
        cfmt = baslik_bicimi(seviye)
    elif stil in ("Title", "Konu Başlığı"):
        bfmt.setProperty(QTextFormat.UserProperty + 1, "baslik")      # belge basligi
        cfmt = baslik_bicimi(0)
    pf = p.paragraph_format
    hiza = p.alignment if p.alignment is not None else pf.alignment
    if hiza is not None:
        bfmt.setAlignment(_qt_hiza(hiza))
    if pf.line_spacing is not None and not hasattr(pf.line_spacing, "pt"):
        bfmt.setLineHeight(float(pf.line_spacing) * 100, QTextBlockFormat.ProportionalHeight.value)
    if pf.space_before is not None:
        bfmt.setTopMargin(pf.space_before.pt * PT_PX)
    if pf.space_after is not None:
        bfmt.setBottomMargin(pf.space_after.pt * PT_PX)
    if pf.left_indent is not None:
        bfmt.setLeftMargin(pf.left_indent.pt * PT_PX)
    if pf.first_line_indent is not None:
        bfmt.setTextIndent(pf.first_line_indent.pt * PT_PX)
    if pf.page_break_before:
        bfmt.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysBefore)
    liste = None
    numpr = p._p.find(qn("w:pPr") + "/" + qn("w:numPr")) if p._p.find(qn("w:pPr")) is not None else None
    if numpr is not None:
        ilvl = numpr.find(qn("w:ilvl")); num_id = numpr.find(qn("w:numId"))
        seviye = int(ilvl.get(qn("w:val"))) if ilvl is not None else 0
        tur = numara.get(num_id.get(qn("w:val")) if num_id is not None else "", {}).get(seviye, "decimal")
        liste = ("madde" if tur == "bullet" else "numara", seviye)
    elif re.match(r"^List (Bullet|Number)( \d)?$", stil):
        s = re.match(r"^List (Bullet|Number)( (\d))?$", stil)
        liste = ("madde" if s.group(1) == "Bullet" else "numara", int(s.group(3) or 1) - 1)
    return bfmt, liste, cfmt


def _qt_hiza(hiza) -> "Qt.Alignment":
    from PySide6.QtCore import Qt
    ad = getattr(hiza, "name", str(hiza))
    return {"LEFT": Qt.AlignLeft, "RIGHT": Qt.AlignRight, "CENTER": Qt.AlignHCenter,
            "JUSTIFY": Qt.AlignJustify, "DISTRIBUTE": Qt.AlignJustify}.get(ad, Qt.AlignLeft)


def _listeye_kat(imlec: QTextCursor, liste) -> None:
    """Ayni tur ve duzeydeki onceki listeye ekle (numara surer); yoksa yeni liste."""
    tur, seviye = liste
    stil = QTextListFormat.ListDisc if tur == "madde" else QTextListFormat.ListDecimal
    if seviye >= 1:
        stil = QTextListFormat.ListCircle if tur == "madde" else QTextListFormat.ListLowerAlpha
    onceki = imlec.block().previous()
    if onceki.isValid() and onceki.textList() is not None:
        lst = onceki.textList()
        if lst.format().style() == stil and lst.format().indent() == seviye + 1:
            lst.add(imlec.block())
            return
    # ayni turdeki son listeyi ara (arada baska duzey varsa numara surer)
    b = onceki
    while b.isValid() and b.textList() is not None:
        lst = b.textList()
        if lst.format().style() == stil and lst.format().indent() == seviye + 1:
            lst.add(imlec.block())
            return
        b = b.previous()
    f = QTextListFormat(); f.setStyle(stil); f.setIndent(seviye + 1)
    imlec.createList(f)


def _run_bicimi(r, temel: QTextCharFormat | None) -> QTextCharFormat:
    f = QTextCharFormat(temel) if temel is not None else QTextCharFormat()
    font = r.font
    if r.bold is not None:
        f.setFontWeight(QFont.Bold if r.bold else QFont.Normal)
    if r.italic is not None:
        f.setFontItalic(bool(r.italic))
    if r.underline:
        f.setFontUnderline(True)
    if font.strike:
        f.setFontStrikeOut(True)
    if font.name:
        f.setFontFamilies([font.name])
    if font.size is not None:
        f.setFontPointSize(font.size.pt)
    if font.color is not None and font.color.type is not None and font.color.rgb is not None:
        f.setForeground(QBrush(QColor("#" + str(font.color.rgb))))
    if font.highlight_color is not None:
        ad = getattr(font.highlight_color, "name", "")
        if ad in VURGU:
            f.setBackground(QBrush(QColor(VURGU[ad])))
    if font.superscript:
        f.setVerticalAlignment(QTextCharFormat.AlignSuperScript)
    elif font.subscript:
        f.setVerticalAlignment(QTextCharFormat.AlignSubScript)
    return f


def _runlari_yaz(imlec, p, d, belge, resim_no, temel) -> bool:
    """Paragrafin runlarini yazar. Donus: paragraf icinde sayfa sonu vardi mi
    (sonraki blok yeni sayfada baslar)."""
    from docx.oxml.ns import qn
    sayfa_sonu = False
    for r_el in p._p.iter(qn("w:r")):
        from docx.text.run import Run
        r = Run(r_el, p)
        bicim = _run_bicimi(r, temel)
        for cocuk in r_el.iterchildren():
            if cocuk.tag == qn("w:t"):
                imlec.insertText(cocuk.text or "", bicim)
            elif cocuk.tag == qn("w:tab"):
                imlec.insertText("\t", bicim)
            elif cocuk.tag == qn("w:br"):
                if cocuk.get(qn("w:type")) == "page":
                    sayfa_sonu = True
                else:
                    imlec.insertText(" ", bicim)            # satir sonu (ayni paragraf)
            elif cocuk.tag == qn("w:drawing"):
                _resim_oku(imlec, cocuk, d, belge, resim_no)
    return sayfa_sonu


def _resim_oku(imlec, cizim, d, belge, resim_no) -> None:
    from docx.oxml.ns import qn
    blip = None
    for e in cizim.iter():
        if e.tag.endswith("}blip"):
            blip = e
            break
    if blip is None:
        return
    rid = blip.get(qn("r:embed"))
    try:
        parca = d.part.related_parts[rid]
    except KeyError:
        return
    img = QImage.fromData(parca.blob)
    if img.isNull():
        return
    resim_no[0] += 1
    ad = f"belge://resim{resim_no[0]}"
    belge.addResource(QTextDocument.ImageResource, QUrl(ad), img)
    f = QTextImageFormat(); f.setName(ad)
    ext = None
    for e in cizim.iter():
        if e.tag.endswith("}extent"):
            ext = e
            break
    if ext is not None:
        f.setWidth(int(ext.get("cx")) / EMU_INC * EKRAN_DPI)
        f.setHeight(int(ext.get("cy")) / EMU_INC * EKRAN_DPI)
    imlec.insertImage(f)


def _tablo_oku(imlec, t, d, belge, resim_no, numara) -> None:
    from docx.oxml.ns import qn
    satirlar = t.rows
    if not satirlar:
        return
    sutun = max(len(r._tr.findall(qn("w:tc"))) and sum(
        int((tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan")).get(qn("w:val"))
             if tc.find(qn("w:tcPr")) is not None and tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan")) is not None else 1))
        for tc in r._tr.findall(qn("w:tc"))) for r in satirlar)
    if imlec.block().text() or imlec.block().previous().isValid():
        imlec.movePosition(QTextCursor.End)
    tf = QTextTableFormat(); tf.setBorder(0.8); tf.setCellPadding(4); tf.setCellSpacing(0)
    tf.setBorderCollapse(True)
    tf.setWidth(QTextLength(QTextLength.PercentageLength, 100))
    tablo = imlec.insertTable(len(satirlar), max(1, sutun), tf)
    from docx.table import _Cell
    for i, r in enumerate(satirlar):
        j = 0
        for tc in r._tr.findall(qn("w:tc")):
            span_el = tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan")) if tc.find(qn("w:tcPr")) is not None else None
            span = int(span_el.get(qn("w:val"))) if span_el is not None else 1
            if j >= tablo.columns():
                break
            hucre = tablo.cellAt(i, j)
            ic = hucre.firstCursorPosition()
            hcell = _Cell(tc, t)
            for k, p in enumerate(hcell.paragraphs):
                bfmt, _, cfmt0 = _paragraf_bicimi(p, numara)
                if k == 0:
                    ic.setBlockFormat(bfmt)
                else:
                    ic.insertBlock(bfmt)
                _runlari_yaz(ic, p, d, belge, resim_no, cfmt0)
            if span > 1:
                tablo.mergeCells(i, j, 1, min(span, tablo.columns() - j))
            j += span


# --------------------------------------------------------------------------
# YAZMA: QTextDocument -> .docx
# --------------------------------------------------------------------------
def docx_yaz(belge: QTextDocument, sayfa: SayfaAyari, yol) -> None:
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.shared import Mm, Pt
    d = docx.Document()
    gövde = d.element.body
    for p in list(d.paragraphs):                # bos sablon paragrafi
        p._p.getparent().remove(p._p)
    s = d.sections[0]
    s.page_width, s.page_height = Mm(sayfa.genislik), Mm(sayfa.yukseklik)
    s.orientation = WD_ORIENT.LANDSCAPE if sayfa.yatay else WD_ORIENT.PORTRAIT
    s.top_margin, s.bottom_margin = Mm(sayfa.ust), Mm(sayfa.alt)
    s.left_margin, s.right_margin = Mm(sayfa.sol), Mm(sayfa.sag)
    vf = belge.defaultFont()
    normal = d.styles["Normal"].font
    normal.name = vf.family()
    if vf.pointSizeF() > 0:
        normal.size = Pt(vf.pointSizeF())
    _ust_alt_bilgi(s, sayfa)
    _cerceve_yaz(d, gövde, belge.rootFrame(), belge)
    d.save(str(yol))


def _ust_alt_bilgi(s, sayfa: SayfaAyari) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    if sayfa.ust_bilgi:
        p = s.header.paragraphs[0]; p.text = sayfa.ust_bilgi; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if sayfa.alt_bilgi or sayfa.sayfa_no:
        p = s.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if sayfa.alt_bilgi:
            p.add_run(sayfa.alt_bilgi + ("  ·  " if sayfa.sayfa_no else ""))
        if sayfa.sayfa_no:
            _alan_ekle(p.add_run(), "PAGE")


def _alan_ekle(run, komut: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for tur, metin in (("begin", None), (None, komut), ("separate", None), (None, "1"), ("end", None)):
        if tur:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tur)
        elif metin == komut:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = f" {komut} "
        else:
            e = OxmlElement("w:t"); e.text = metin
        run._r.append(e)


def _cerceve_yaz(d, kap, cerceve: QTextFrame, belge) -> None:
    from PySide6.QtGui import QTextTable
    it = cerceve.begin()
    while not it.atEnd():
        alt = it.currentFrame()
        blok = it.currentBlock()
        if alt is not None and isinstance(alt, QTextTable):
            _tablo_yaz(d, kap, alt, belge)
        elif blok.isValid():
            p = _yeni_paragraf(d, kap)
            _blok_yaz(d, p, blok, belge)
        it += 1


def _yeni_paragraf(d, kap):
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    p_el = OxmlElement("w:p")
    sect = kap.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr")
    if sect is not None:
        sect.addprevious(p_el)
    else:
        kap.append(p_el)
    return Paragraph(p_el, d._body if hasattr(d, "_body") else d)


def _blok_yaz(d, p, blok, belge) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    from PySide6.QtCore import Qt
    bf = blok.blockFormat()
    seviye = bf.headingLevel()
    liste = blok.textList()
    if seviye:
        p.style = d.styles[BASLIK_STILI.get(min(seviye, 3), "Heading 3")]
    elif bf.property(QTextFormat.UserProperty + 1) == "baslik":
        p.style = d.styles["Title"]
    elif liste is not None:
        lf = liste.format()
        madde = lf.style() in (QTextListFormat.ListDisc, QTextListFormat.ListCircle, QTextListFormat.ListSquare)
        duzey = max(1, min(lf.indent(), 3))
        ad = ("List Bullet" if madde else "List Number") + ("" if duzey == 1 else f" {duzey}")
        p.style = d.styles[ad]
    hiza = bf.alignment()
    if hiza & Qt.AlignHCenter:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif hiza & Qt.AlignRight:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif hiza & Qt.AlignJustify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    if bf.lineHeightType() == QTextBlockFormat.ProportionalHeight.value and bf.lineHeight() > 0:
        pf.line_spacing = round(bf.lineHeight() / 100.0, 2)
    if bf.topMargin():
        pf.space_before = Pt(bf.topMargin() / PT_PX)
    if bf.bottomMargin():
        pf.space_after = Pt(bf.bottomMargin() / PT_PX)
    if bf.leftMargin() and liste is None:
        pf.left_indent = Pt(bf.leftMargin() / PT_PX)
    if bf.textIndent():
        pf.first_line_indent = Pt(bf.textIndent() / PT_PX)
    if bf.pageBreakPolicy() & QTextFormat.PageBreak_AlwaysBefore:
        pf.page_break_before = True
    it = blok.begin()
    while not it.atEnd():
        parca = it.fragment()
        if parca.isValid():
            cf = parca.charFormat()
            if cf.isImageFormat():
                _resim_yaz(p, cf.toImageFormat(), belge)
            else:
                metin = parca.text()
                satirlar = metin.split(" ")
                for k, s in enumerate(satirlar):
                    if k:
                        p.add_run().add_break()
                    if s:
                        _run_yaz(p.add_run(s.replace("￼", "")), cf, seviye or bf.property(QTextFormat.UserProperty + 1))
        it += 1


def _run_yaz(r, cf: QTextCharFormat, baslik) -> None:
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Pt, RGBColor
    if cf.fontWeight() >= QFont.Bold and not baslik:
        r.bold = True
    if cf.fontItalic():
        r.italic = True
    if cf.fontUnderline():
        r.underline = True
    if cf.fontStrikeOut():
        r.font.strike = True
    aileler = cf.fontFamilies()
    if aileler:
        r.font.name = aileler[0]
    if cf.fontPointSize() > 0 and not baslik:
        r.font.size = Pt(cf.fontPointSize())
    if cf.hasProperty(QTextFormat.ForegroundBrush):
        c = cf.foreground().color()
        r.font.color.rgb = RGBColor(c.red(), c.green(), c.blue())
    if cf.hasProperty(QTextFormat.BackgroundBrush) and cf.background().style() != 0:
        r.font.highlight_color = getattr(WD_COLOR_INDEX, en_yakin_vurgu(cf.background().color()))
    if cf.verticalAlignment() == QTextCharFormat.AlignSuperScript:
        r.font.superscript = True
    elif cf.verticalAlignment() == QTextCharFormat.AlignSubScript:
        r.font.subscript = True


def _resim_yaz(p, imf: QTextImageFormat, belge) -> None:
    from docx.shared import Inches
    kaynak = belge.resource(QTextDocument.ImageResource, QUrl(imf.name()))
    img = kaynak if isinstance(kaynak, QImage) else QImage(kaynak) if kaynak is not None else QImage(imf.name())
    if img.isNull():
        return
    ba = QByteArray(); tampon = QBuffer(ba); tampon.open(QIODevice.WriteOnly)
    img.save(tampon, "PNG")
    genislik = imf.width() if imf.width() > 0 else img.width()
    p.add_run().add_picture(io.BytesIO(bytes(ba.data())), width=Inches(genislik / EKRAN_DPI))


def _tablo_yaz(d, kap, tablo, belge) -> None:
    from docx.oxml import OxmlElement
    from docx.table import Table
    satir, sutun = tablo.rows(), tablo.columns()
    t = d.add_table(rows=satir, cols=sutun)            # sona eklenir; dogru yere tasinir
    t.style = d.styles["Table Grid"]
    sect = kap.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr")
    tbl = t._tbl
    tbl.getparent().remove(tbl)
    if sect is not None:
        sect.addprevious(tbl)
    else:
        kap.append(tbl)
    t = Table(tbl, d._body if hasattr(d, "_body") else d)
    yapildi = set()
    for i in range(satir):
        for j in range(sutun):
            hucre = tablo.cellAt(i, j)
            if (hucre.row(), hucre.column()) in yapildi:
                continue
            yapildi.add((hucre.row(), hucre.column()))
            hedef = t.cell(i, j)
            if hucre.columnSpan() > 1 or hucre.rowSpan() > 1:
                hedef = hedef.merge(t.cell(i + hucre.rowSpan() - 1, j + hucre.columnSpan() - 1))
            ilk = True
            blok = hucre.firstCursorPosition().block()
            son = hucre.lastCursorPosition().block()
            while blok.isValid():
                p = hedef.paragraphs[0] if ilk else hedef.add_paragraph()
                ilk = False
                _blok_yaz(d, p, blok, belge)
                if blok == son:
                    break
                blok = blok.next()
